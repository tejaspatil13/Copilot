"""Write results (a single search, or an employee's whole sheet) into the client's spreadsheet template."""
import re
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models import ResultRecord

# A-E are the client's original columns (B = the address as the business lists it). F onwards mirror
# the dashboard. "comments" and "Status" are left empty for the caller to fill in.
HEADERS = (
    ("business name", None), ("add", None), ("RV", None), ("Contact no", None), ("comments", None),
    ("Status", 16), ("Other numbers at this address", 48), ("Note", 40), ("Description (VO)", 28),
    ("Postcode", 10), ("Google Maps", 13), ("Gov.uk record", 14), ("VO address", 45), ("VO UARN", 14),
    ("Search", 30), ("Added", 17),
)
COL_PHONE, COL_MAPS, COL_GOV = 4, 11, 12
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")
_LINK = Font(name="Calibri", size=11, color="0563C1", underline="single")
_FILENAME_BAD = re.compile(r"[^A-Za-z0-9 ._()&'-]+")


@dataclass
class ExportRow:
    record: ResultRecord
    search: str = ""  # which search the row came from, e.g. "LONDON · EC1V 9LT · CO"
    added: datetime | None = None


def safe_cell(value):
    """Neutralise spreadsheet formula injection for any text that came from an external source."""
    if isinstance(value, str):
        value = value.replace("\x00", "")
        if value.startswith(_FORMULA_START):
            return "'" + value
    return value


def export_filename(town: str, postcode: str, when: datetime) -> str:
    """Default download name, e.g. UK-Businesses_Coventry_CV1-1FJ_2026-10-26_2315.xlsx"""
    parts = [p for p in (town, postcode) if p] or ["search"]
    slug = "_".join(re.sub(r"[^A-Za-z0-9]+", "-", p).strip("-") for p in parts)[:60]
    return f"UK-Businesses_{slug}_{when:%Y-%m-%d_%H%M}.xlsx"


def clean_filename(name: str | None, fallback: str) -> str:
    """User-chosen download name -> safe "<name>.xlsx" (no paths, no odd characters)."""
    base = _FILENAME_BAD.sub(" ", re.sub(r"\.xlsx$", "", (name or "").strip(), flags=re.I))
    base = re.sub(r"\s+", " ", base).strip(" .")[:80]
    return f"{base}.xlsx" if base else fallback


ATTRIBUTION = (
    ("Rateable values & addresses",
     "Valuation Office (HMRC) compiled rating list, © Crown copyright. Supplied under the VO rating list "
     "terms: use is restricted to non-domestic rating purposes; onward disclosure only to parties using "
     "it for those purposes; delete when no longer needed."),
    ("Business names & phone numbers",
     "Overture Maps Foundation places (CDLA-Permissive-2.0; includes data from OpenStreetMap contributors, "
     "Meta, Microsoft, Foursquare and others — see https://docs.overturemaps.org/attribution/)."),
    ("Several businesses at one address",
     "Shared buildings list several businesses at the same address. The most likely one's number is in "
     "'Contact no'; the others are in 'Other numbers at this address'. Use the Google Maps link to check."),
    ("Before calling",
     "Unsolicited sales/marketing calls must be screened against TPS and CTPS (Privacy and Electronic "
     "Communications Regulations). This file has NOT been screened."),
)


def write_workbook(
    template: Path, out_path: Path, rows: list[ExportRow], rv_header: str,
    about: list[tuple[str, str]] | None = None,
) -> None:
    wb = openpyxl.load_workbook(template)
    ws = wb.active
    style_row = [copy(c._style) for c in ws[2][:5]]
    header_style = copy(ws["A1"]._style)

    # Always open at the top-left (the client's original workbook was saved scrolled to row ~2,870).
    ws.sheet_view.topLeftCell = "A1"
    for sel in ws.sheet_view.selection:
        sel.activeCell, sel.sqref = "A2", "A2"
    for idx in [i for i in ws.row_dimensions if i > 1]:
        del ws.row_dimensions[idx]

    for i, (title, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i, value=rv_header if i == 3 else title)
        if i > 5:
            cell._style = copy(header_style)
            ws.column_dimensions[get_column_letter(i)].width = width

    for r, row in enumerate(rows, start=2):
        rec = row.record
        added = row.added or rec.collected_at
        values = (
            rec.business_name, rec.business_address or rec.address, rec.rv, rec.phone, "",
            "", "; ".join(rec.other_numbers), rec.match_note, rec.description, rec.postcode,
            "Open map", "Open record", rec.address, rec.uarn, row.search,
            added.strftime("%Y-%m-%d %H:%M") if added else "",
        )
        for c, value in enumerate(values, start=1):
            value = "" if value is None else value
            # The phone is our own validated "+44 ..." string, never a formula; keep the "+" visible.
            cell = ws.cell(row=r, column=c, value=value if c == COL_PHONE else safe_cell(value))
            if c <= 5:
                cell._style = copy(style_row[c - 1])
        ws.cell(row=r, column=COL_PHONE).number_format = "@"
        for col, url in ((COL_MAPS, rec.google_maps_url), (COL_GOV, rec.source_url)):
            link = ws.cell(row=r, column=col)
            link.hyperlink, link.font = url, _LINK

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(len(rows) + 1, 2)}"
    ws.freeze_panes = "A2"

    if about:
        info = wb.create_sheet("About")
        info.column_dimensions["A"].width = 28
        info.column_dimensions["B"].width = 110
        for i, (key, value) in enumerate(about, start=1):
            info.cell(row=i, column=1, value=safe_cell(key)).font = Font(bold=True)
            info.cell(row=i, column=2, value=safe_cell(value))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(".tmp")
    wb.save(tmp)
    tmp.replace(out_path)
