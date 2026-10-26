"""End-to-end API tests with a mocked contact provider (no network)."""
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.contacts.base import ContactCandidate, ProviderRateLimited


def download_name(response) -> str:
    """File name from Content-Disposition (plain or RFC 5987 'filename*=utf-8\'\'...' form)."""
    from urllib.parse import unquote

    cd = response.headers["content-disposition"]
    if "filename*=" in cd:
        return unquote(cd.split("filename*=", 1)[1].split("''", 1)[1])
    return cd.split('filename="', 1)[1].rstrip('"')


class FakeProvider:
    name = "fake"

    def __init__(self, data: dict[str, list[ContactCandidate]], rate_limited: set[str] = frozenset()):
        self.data, self.rate_limited, self.calls = data, set(rate_limited), []

    async def candidates(self, record):
        self.calls.append(record.uarn)
        if record.uarn in self.rate_limited:
            raise ProviderRateLimited(retry_after=1)
        return self.data.get(record.uarn, [])


GW = ContactCandidate(name="Games Workshop", address="13 Trinity St, Coventry", postcode="CV1 1FJ",
                      phone="+44 24 7622 7311", source="fake", source_url="https://example.com/gw")


@pytest.fixture
def make_client(settings, monkeypatch):
    def factory(provider=None):
        import app.main as main
        monkeypatch.setattr(main, "settings", settings)
        monkeypatch.setattr(main, "build_providers", lambda s: [provider] if provider else [])
        return TestClient(main.app)
    return factory


def wait(client, job_id, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["state"] not in ("queued", "running"):
            return job
        time.sleep(0.05)
    raise AssertionError(f"job did not finish: {job}")


BIG = ContactCandidate(name="Big Shop", address="9-11 Trinity Street, Coventry", postcode="CV1 1FJ",
                       phone="024 7600 1111", source="fake")


def test_full_search_excel_has_only_properties_with_phone(make_client):
    provider = FakeProvider({"300675222": [GW], "235950222": [BIG]})
    with make_client(provider) as client:
        assert client.get("/api/health").json()["status"] == "ok"
        assert len(client.get("/api/meta").json()["description_codes"]) == 117
        r = client.post("/api/jobs", json={"town": "Coventry", "postcode": "CV1 1FJ", "rv_from": 12000})
        assert r.status_code == 202
        job = wait(client, r.json()["id"])
        assert job["state"] == "completed", job
        s = job["summary"]
        assert (s["source_records"], s["phones_found"], s["verified"], s["no_contact"]) == (4, 2, 2, 2)
        assert "2 have a phone number" in job["message"]

        # The dashboard still shows every listed property with its status...
        assert client.get(f"/api/jobs/{job['id']}/records").json()["total"] == 4
        items = client.get(f"/api/jobs/{job['id']}/records?has_phone=true").json()["items"]
        assert {i["business_name"] for i in items} == {"Games Workshop", "Big Shop"}

        # ...but the spreadsheet only contains the ones with a phone number.
        dl = client.get(f"/api/jobs/{job['id']}/download")
        assert dl.status_code == 200
        assert "UK-Businesses_Coventry_CV1-1FJ_" in dl.headers["content-disposition"]
        assert dl.headers["x-content-type-options"] == "nosniff"
        ws = openpyxl.load_workbook(client.app.state.jobs.settings.export_dir / f"{job['id']}.xlsx")["Results"]
        assert ws.max_row == 3
        assert sorted(ws.cell(r, 4).value for r in (2, 3)) == ["+44 24 7600 1111", "+44 24 7622 7311"]


def test_no_phone_numbers_means_no_spreadsheet(make_client):
    with make_client(FakeProvider({})) as client:
        job = wait(client, client.post("/api/jobs", json={"postcode": "CV1 1FJ"}).json()["id"])
        assert job["state"] == "no_results" and not job["download_ready"]
        assert "none has a reliable phone number" in job["message"]


def test_rate_limited_lookup_can_be_retried_and_only_reruns_failures(make_client):
    provider = FakeProvider({"300675222": [GW]}, rate_limited={"300675222"})
    with make_client(provider) as client:
        job = wait(client, client.post("/api/jobs", json={"postcode": "CV1 1FJ"}).json()["id"])
        assert job["summary"]["lookup_failed"] == 1 and "Retry" in job["message"]

        provider.rate_limited.clear()
        provider.calls.clear()
        client.post(f"/api/jobs/{job['id']}/retry")
        job = wait(client, job["id"])
        assert job["state"] == "completed" and job["summary"]["verified"] == 1 and job["download_ready"]
        assert provider.calls == ["300675222"]


def test_missing_phone_dataset_is_reported(make_client):
    with make_client(None) as client:
        job = wait(client, client.post("/api/jobs", json={"postcode": "CV1 1FJ"}).json()["id"])
        assert job["state"] == "failed"
        assert "overture_import" in job["message"]


def test_no_results_gives_message_and_no_file(make_client):
    with make_client(None) as client:
        job = wait(client, client.post("/api/jobs", json={"town": "Atlantis"}).json()["id"])
        assert job["state"] == "no_results" and "Atlantis" in job["message"]
        assert not job["download_ready"]
        assert client.get(f"/api/jobs/{job['id']}/download").status_code == 404


def test_validation_errors_are_readable(make_client):
    with make_client(None) as client:
        r = client.post("/api/jobs", json={"postcode": "nonsense"})
        assert r.status_code == 422
        assert r.json()["detail"] == ["Postcode: must be a full UK postcode, sector (e.g. 'CV1 1') or outward code (e.g. 'CV1')"]
        assert client.post("/api/jobs", json={}).json()["detail"] == ["Enter a town, a postcode, or both."]


@pytest.mark.parametrize("job_id", ["../../etc/passwd", "..%2F..%2Fapp.sqlite", "0" * 32, "abc"])
def test_unknown_or_malicious_job_ids_404(make_client, job_id):
    with make_client(None) as client:
        assert client.get(f"/api/jobs/{job_id}/download").status_code == 404


def test_password_protection(make_client, settings):
    settings.app_password = "s3cret"
    with make_client(None) as client:
        assert client.get("/api/jobs").status_code == 401
        assert client.get("/api/health").status_code == 200
        assert client.get("/api/jobs", auth=("user", "s3cret")).status_code == 200
        assert client.get("/api/jobs", auth=("user", "wrong")).status_code == 401


def test_damaged_search_history_is_set_aside_and_app_still_works(make_client, settings):
    settings.app_db_path.write_bytes(b"this is not a sqlite database" * 100)
    with make_client(FakeProvider({"300675222": [GW]})) as client:
        job = wait(client, client.post("/api/jobs", json={"postcode": "CV1 1FJ"}).json()["id"])
        assert job["state"] == "completed"
    assert list(settings.data_dir.glob("app.sqlite.corrupt-*"))


def test_download_uses_the_name_the_user_chose(make_client):
    with make_client(FakeProvider({"300675222": [GW]})) as client:
        job = wait(client, client.post("/api/jobs", json={"postcode": "CV1 1FJ"}).json()["id"])
        dl = client.get(f"/api/jobs/{job['id']}/download", params={"filename": "Sara Coventry list"})
        assert dl.status_code == 200 and download_name(dl) == "Sara Coventry list.xlsx"
        dl = client.get(f"/api/jobs/{job['id']}/download", params={"filename": "../../secret"})
        assert download_name(dl) == "secret.xlsx"


def test_history_is_per_employee_and_searches_can_be_deleted(make_client, settings):
    with make_client(FakeProvider({"300675222": [GW]})) as client:
        a = wait(client, client.post("/api/jobs?employee=Sara", json={"postcode": "CV1 1FJ"}).json()["id"])
        wait(client, client.post("/api/jobs?employee=Tom", json={"postcode": "CV1 1FJ"}).json()["id"])
        assert a["employee"] == "Sara" and a["label"] == "CV1 1FJ"
        assert [j["id"] for j in client.get("/api/jobs?employee=sara").json()] == [a["id"]]

        assert client.delete(f"/api/jobs/{a['id']}").status_code == 204
        assert client.get(f"/api/jobs/{a['id']}").status_code == 404
        assert not (settings.export_dir / f"{a['id']}.xlsx").exists()
        assert client.get("/api/jobs?employee=Sara").json() == []
        assert client.post("/api/jobs?employee=<script>", json={"postcode": "CV1 1FJ"}).status_code == 422


def test_my_sheet_accumulates_searches_and_downloads_as_one_file(make_client, tmp_path):
    provider = FakeProvider({"300675222": [GW], "235950222": [BIG], "999000001": [ContactCandidate(
        name="Dhabba", address="1 Lower Holyhead Rd, Coventry", postcode="CV1 3AX", phone="024 7655 9900",
        source="fake")]})
    with make_client(provider) as client:
        one = wait(client, client.post("/api/jobs?employee=Sara", json={"postcode": "CV1 1FJ"}).json()["id"])
        two = wait(client, client.post("/api/jobs?employee=Sara", json={"postcode": "CV1 3AX"}).json()["id"])

        r = client.post("/api/sheet/add", json={"employee": "Sara", "job_id": one["id"]}).json()
        assert r == {"added": 2, "already_in_sheet": 0, "total": 2}
        r = client.post("/api/sheet/add", json={"employee": "sara ", "job_id": two["id"]}).json()
        assert r == {"added": 1, "already_in_sheet": 0, "total": 3}
        # adding the same search again doesn't duplicate rows
        r = client.post("/api/sheet/add", json={"employee": "Sara", "job_id": one["id"]}).json()
        assert r == {"added": 0, "already_in_sheet": 2, "total": 3}
        # another employee's sheet is separate
        assert client.get("/api/sheet", params={"employee": "Tom"}).json()["total"] == 0

        sheet = client.get("/api/sheet", params={"employee": "Sara"}).json()
        assert {i["search"] for i in sheet["items"]} == {"CV1 1FJ", "CV1 3AX"}

        dl = client.post("/api/sheet/download", json={"employee": "Sara", "filename": "Sara - October"})
        assert dl.status_code == 200 and download_name(dl) == "Sara - October.xlsx"
        f = tmp_path / "sheet.xlsx"
        f.write_bytes(dl.content)
        ws = openpyxl.load_workbook(f)["Results"]
        assert ws.max_row == 4 and ws.sheet_view.topLeftCell == "A1"

        # filtered download: only the rows the user is looking at
        dl = client.post("/api/sheet/download", json={"employee": "Sara", "uarns": ["999000001"]})
        f.write_bytes(dl.content)
        assert openpyxl.load_workbook(f)["Results"].max_row == 2
        assert download_name(dl).startswith("Sara - My sheet - ")   # default name uses the employee

        r = client.post("/api/sheet/remove", json={"employee": "Sara", "uarns": ["999000001"]}).json()
        assert r == {"removed": 1, "total": 2}
        assert client.delete("/api/sheet", params={"employee": "Sara"}).json()["total"] == 0
        assert client.post("/api/sheet/download", json={"employee": "Sara"}).status_code == 404
