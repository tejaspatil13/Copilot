from datetime import datetime, timezone

import openpyxl
import pytest

from app.excel import ExportRow, clean_filename, export_filename, safe_cell, write_workbook
from app.models import ContactStatus, ResultRecord
from tests.conftest import TEMPLATE


def rec(**kw):
    base = dict(
        uarn="300675222", address="13, Trinity Street, Coventry, CV1 1FJ", street="Trinity Street",
        number_or_name="13", town="Coventry", postcode="CV1 1FJ", description="Shop and Premises", rv=22000,
        business_name="Games Workshop", phone="+44 24 7622 7311", contact_status=ContactStatus.VERIFIED,
        match_score=95, match_reasons=["postcode matches"], contact_source="overture",
        contact_source_url="https://example.com/gw", collected_at=datetime(2026, 10, 26, tzinfo=timezone.utc),
    )
    return ResultRecord(**(base | kw))


@pytest.mark.parametrize("value", ["=HYPERLINK(\"x\")", "+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"])
def test_safe_cell_neutralises_formulas(value):
    assert safe_cell(value).startswith("'")


def test_safe_cell_leaves_normal_values():
    assert safe_cell("Games Workshop") == "Games Workshop"
    assert safe_cell(22000) == 22000


def test_export_filename_is_safe():
    name = export_filename("../../etc/passwd", "CV1 1FJ", datetime(2026, 10, 26, 23, 15))
    assert name == "UK-Businesses_etc-passwd_CV1-1FJ_2026-10-26_2315.xlsx"
    assert "/" not in export_filename("a/b\\c", "", datetime(2026, 1, 1))


def test_clean_filename():
    assert clean_filename("Sara - My sheet", "x.xlsx") == "Sara - My sheet.xlsx"
    assert clean_filename("../../etc/passwd.xlsx", "x.xlsx") == "etc passwd.xlsx"
    assert clean_filename("   ", "default.xlsx") == "default.xlsx"
    assert clean_filename(None, "default.xlsx") == "default.xlsx"


def test_workbook_matches_template_and_opens_at_the_top(tmp_path):
    out = tmp_path / "out.xlsx"
    rows = [
        ExportRow(rec(other_numbers=["Warhammer: +44 24 7622 7311"], match_note="Same building number and postcode"),
                  "COVENTRY · CV1 1FJ"),
        ExportRow(rec(uarn="1", business_name="=cmd|' /C calc'!A0", phone="")),
    ]
    write_workbook(TEMPLATE, out, rows, "RV (2026 list)")

    tpl = openpyxl.load_workbook(TEMPLATE).active
    ws = openpyxl.load_workbook(out).active
    # Opens at the top (the client's sample was saved scrolled to row ~2,870)
    assert ws.sheet_view.topLeftCell == "A1"
    assert not [i for i in ws.row_dimensions if i > 1 and ws.row_dimensions[i].height]
    # Template columns A, B, D, E keep their exact headers; C carries the RV label; F.. mirror the dashboard
    assert [ws.cell(1, c).value for c in (1, 2, 4, 5)] == [tpl.cell(1, c).value for c in (1, 2, 4, 5)]
    assert ws["C1"].value == "RV (2026 list)"
    assert [ws.cell(1, c).value for c in range(6, 13)] == [
        "Status", "Other numbers at this address", "Note", "Description (VO)", "Postcode", "Google Maps",
        "Gov.uk record"]
    assert ws.cell(1, 15).value == "Search" and ws.cell(1, 16).value == "Added"
    # Formatting carried over from the template's data row
    assert ws["C2"].number_format == tpl["C2"].number_format == "#,##0"
    assert ws.column_dimensions["B"].width == tpl.column_dimensions["B"].width
    assert ws["A1"].font.b
    # Values
    assert ws["A2"].value == "Games Workshop" and ws["C2"].value == 22000
    assert ws["D2"].value == "+44 24 7622 7311" and ws["D2"].data_type == "s"   # "+" kept, not escaped
    assert ws["B2"].value == "13, Trinity Street, Coventry, CV1 1FJ"            # VO address when no listed one
    assert ws["E2"].value is None and ws["F2"].value is None                     # comments + Status left empty
    assert ws["G2"].value == "Warhammer: +44 24 7622 7311"
    assert ws["H2"].value == "Same building number and postcode"
    assert ws["O2"].value == "COVENTRY · CV1 1FJ"
    assert ws["L2"].hyperlink.target.endswith("/valuations/start/300675222")
    assert ws["K2"].hyperlink.target == (
        "https://www.google.com/maps/search/?api=1&query=Games+Workshop%2C+13%2C+Trinity+Street%2C+Coventry%2C+CV1+1FJ")
    assert ws["A3"].value.startswith("'=")           # formula injection neutralised
    assert ws.auto_filter.ref.startswith("A1:")
    assert len(ws["A"]) == 3


def test_listed_business_address_goes_in_column_b(tmp_path):
    out = tmp_path / "out.xlsx"
    write_workbook(TEMPLATE, out, [ExportRow(rec(business_address="13 Trinity St, Coventry, CV1 1FJ"))], "RV")
    ws = openpyxl.load_workbook(out).active
    assert ws["B2"].value == "13 Trinity St, Coventry, CV1 1FJ" and ws["M2"].value == "13, Trinity Street, Coventry, CV1 1FJ"


def test_google_maps_link_uses_business_listing():
    assert rec(business_name="").google_maps_url.endswith("query=13%2C+Trinity+Street%2C+Coventry%2C+CV1+1FJ")
    assert rec(business_address="13 Trinity St, Coventry, CV1 1FJ").google_maps_url.endswith(
        "query=Games+Workshop%2C+13+Trinity+St%2C+Coventry%2C+CV1+1FJ")


def test_about_sheet_is_added_after_results(tmp_path):
    out = tmp_path / "out.xlsx"
    write_workbook(TEMPLATE, out, [ExportRow(rec())], "RV", about=[("Search", "town=Coventry"), ("Note", "=1+1")])
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Results", "About"]
    assert wb["About"]["B1"].value == "town=Coventry" and wb["About"]["B2"].value == "'=1+1"
